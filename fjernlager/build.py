"""
Samler fjernlager fra én eller flere leverandører til Matrixify-filer for .dk og .com.

Flow:
  1. Hent hver aktiv leverandørs fil (FTP, FTPS eller SFTP)
  2. Oversæt leverandørens format til (stregkode, antal) via en adapter
  3. Læg antal sammen pr. stregkode på tværs af leverandører
  4. Varer der er forsvundet fra listerne (set inden for STATE_DAYS dage) får antal 0
  5. Værn: mangler en fil, er den for gammel, eller er den skrumpet mere end
     max_drop_pct i forhold til sidst, publiceres INTET, og kørslen fejler
  6. Skriv én xlsx pr. butik og upload til Hetzner (+ arkivkopi)

Kørsel i GitHub Actions:   python build.py
Lokal test uden netværk:   python build.py --local neilpryde=C:\\sti\\fil.csv --out C:\\sti\\ud --state C:\\sti\\state.json
"""

import argparse, datetime as dt, io, json, os, sys, tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
STATE_DAYS = 14
NOW = dt.datetime.now(dt.timezone.utc)

# Repo'et er offentligt, og det er Actions-loggen også. Derfor logges antal varer,
# lagertal og filstier kun ved lokal test eller med FJERNLAGER_VERBOSE=1.
VERBOSE = os.environ.get("FJERNLAGER_VERBOSE") == "1"


def vprint(*a):
    if VERBOSE:
        print(*a)


# ---------------------------------------------------------------- adaptere
# En adapter tager en lokal filsti og returnerer [(stregkode, antal), ...].
# Ny leverandør = ny funktion her + en linje i config.json.

def adapter_matrixify_xlsx(path, opts):
    """Fil der allerede er i Matrixify-format (fx den nuværende oos_fjernlager.xlsx)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    i_bc = header.index("Variant Barcode")
    i_qty = next(i for i, h in enumerate(header) if h.startswith("Inventory Available"))
    return [(r[i_bc], r[i_qty]) for r in rows]


def adapter_neilpryde(path, opts):
    """NeilPryde / Wavos. Bygges færdig, når formatet er kendt."""
    raise NotImplementedError("NeilPryde-adapteren mangler - send en eksempelfil")


ADAPTERS = {
    "matrixify_xlsx": adapter_matrixify_xlsx,
    "neilpryde": adapter_neilpryde,
}


# ---------------------------------------------------------------- transport
def secret(name):
    v = os.environ.get(name, "")
    if not v:
        raise SystemExit("Mangler secret %s" % name)
    return v


def fetch(src, dest):
    """Henter src['path'] til dest. Returnerer filens ændringstid (UTC) eller None."""
    kind = src["type"]
    host, user, pw = secret(src["host_secret"]), secret(src["user_secret"]), secret(src["pass_secret"])
    if kind in ("ftp", "ftps"):
        import ftplib
        ftp = ftplib.FTP_TLS(host, timeout=60) if kind == "ftps" else ftplib.FTP(host, timeout=60)
        ftp.login(user, pw)
        if kind == "ftps":
            ftp.prot_p()
        mtime = None
        try:
            resp = ftp.sendcmd("MDTM " + src["path"])          # "213 20260921040012"
            mtime = dt.datetime.strptime(resp.split()[1][:14], "%Y%m%d%H%M%S").replace(tzinfo=dt.timezone.utc)
        except Exception:
            pass
        with open(dest, "wb") as f:
            ftp.retrbinary("RETR " + src["path"], f.write)
        ftp.quit()
        return mtime
    if kind == "sftp":
        import paramiko
        t = paramiko.Transport((host, int(src.get("port", 22))))
        t.connect(username=user, password=pw)
        s = paramiko.SFTPClient.from_transport(t)
        mtime = dt.datetime.fromtimestamp(s.stat(src["path"]).st_mtime, dt.timezone.utc)
        s.get(src["path"], dest)
        s.close(); t.close()
        return mtime
    raise SystemExit("Ukendt transporttype: %s" % kind)


def upload(files, target):
    import paramiko
    t = paramiko.Transport((secret(target["host_secret"]), int(target.get("port", 22))))
    t.connect(username=secret(target["user_secret"]), password=secret(target["pass_secret"]))
    s = paramiko.SFTPClient.from_transport(t)
    for d in sorted({r.rsplit("/", 1)[0] for _, r in files}):
        try:
            s.mkdir(d)
        except IOError:
            pass                          # findes allerede
    for local, remote in files:
        tmp = remote + ".uploading"
        s.put(local, tmp)
        try:
            s.remove(remote)
        except IOError:
            pass
        s.rename(tmp, remote)            # atomisk skift: Matrixify ser aldrig en halv fil
        vprint("  uploadet ->", remote)
    s.close(); t.close()


def read_remote_state(target, path):
    try:
        import paramiko
        t = paramiko.Transport((secret(target["host_secret"]), int(target.get("port", 22))))
        t.connect(username=secret(target["user_secret"]), password=secret(target["pass_secret"]))
        s = paramiko.SFTPClient.from_transport(t)
        with s.open(path) as f:
            data = json.loads(f.read().decode("utf-8"))
        s.close(); t.close()
        return data
    except Exception:
        return {}


# ---------------------------------------------------------------- normalisering
def clean_barcode(v):
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0"):                 # Excel gemmer tal som 4045533780902.0
        s = s[:-2]
    # Shopify-stregkoder må indeholde bogstaver, så kun tomme og "None"-agtige afvises
    if not s or s.lower() in ("none", "null", "nan", "-") or " " in s:
        return None
    return s


def clean_qty(v):
    try:
        q = int(float(str(v).replace(",", ".").strip()))
    except (TypeError, ValueError):
        return 0
    return max(q, 0)


def write_xlsx(path, column, stock):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Variant Barcode", column])
    for bc in sorted(stock):
        ws.append([bc, stock[bc]])
    wb.save(path)


# ---------------------------------------------------------------- hovedforløb
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--local", action="append", default=[], help="navn=sti - brug lokal fil i stedet for at hente")
    ap.add_argument("--out", help="lokal outputmappe (ingen upload)")
    ap.add_argument("--state", help="lokal state-fil (kun ved lokal test)")
    args = ap.parse_args()

    global VERBOSE
    if args.out:
        VERBOSE = True                    # lokal test: vis alt

    cfg = json.load(io.open(os.path.join(HERE, "config.json"), encoding="utf-8"))
    local_files = dict(x.split("=", 1) for x in args.local)
    local_mode = bool(args.out)
    target = cfg["target"]
    state_remote = target["dir"] + "/" + cfg["state_file"]

    if local_mode:
        state = json.load(open(args.state)) if args.state and os.path.exists(args.state) else {}
    else:
        state = read_remote_state(target, state_remote)
    state.setdefault("seen", {})
    state.setdefault("rows", {})

    work = tempfile.mkdtemp()
    merged, per_supplier, errors = {}, {}, []

    for sup in cfg["suppliers"]:
        if not sup.get("enabled"):
            continue
        name = sup["name"]
        try:
            if name in local_files:
                path, mtime = local_files[name], None
            else:
                path = os.path.join(work, name + os.path.splitext(sup["source"]["path"])[1])
                mtime = fetch(sup["source"], path)

            if mtime is not None:
                age_h = (NOW - mtime).total_seconds() / 3600
                vprint("[%s] fil ændret %s (%.0f timer siden)" % (name, mtime.isoformat(), age_h))
                if age_h > sup.get("max_age_hours", 30):
                    errors.append("%s: filen er %.0f timer gammel" % (name, age_h))
                    continue

            rows = ADAPTERS[sup["adapter"]](path, sup.get("options", {}))
            stock, skipped = {}, 0
            for bc, qty in rows:
                b = clean_barcode(bc)
                if b is None:
                    skipped += 1
                    continue
                # Samme stregkode to gange i ÉN fil: tag den højeste, læg ikke sammen.
                # (Mellem leverandører lægges der sammen - se nedenfor.)
                stock[b] = max(stock.get(b, 0), clean_qty(qty))

            prev = state["rows"].get(name)
            vprint("[%s] %d stregkoder (%d uden gyldig stregkode sprunget over), %d stk. på lager, sidst %s"
                  % (name, len(stock), skipped, sum(stock.values()), prev))
            if not stock:
                errors.append("%s: filen gav ingen rækker" % name)
                continue
            if prev and len(stock) < prev * (1 - cfg["max_drop_pct"] / 100.0):
                fald = 100 - len(stock) * 100 // prev
                errors.append("%s: filen er skrumpet %d%% siden sidst (grænse %d%%)"
                              % (name, fald, cfg["max_drop_pct"]))
                continue

            per_supplier[name] = len(stock)
            for b, q in stock.items():
                merged[b] = merged.get(b, 0) + q
        except Exception as e:
            errors.append("%s: %s" % (name, e))

    if errors:
        print("\nINTET PUBLICERET - den forrige fil bliver liggende:")
        for e in errors:
            print("  ::error::" + e)
        sys.exit(1)

    # Varer set inden for STATE_DAYS dage, men ikke med i dag -> 0
    cutoff = (NOW - dt.timedelta(days=STATE_DAYS)).isoformat()
    for b in merged:
        state["seen"][b] = NOW.isoformat()
    zeroed = 0
    for b, last in list(state["seen"].items()):
        if b in merged:
            continue
        if last < cutoff:
            del state["seen"][b]
        else:
            merged[b] = 0
            zeroed += 1
    state["rows"].update(per_supplier)
    state["last_run"] = NOW.isoformat()

    vprint("\nSamlet: %d stregkoder, %d stk., heraf %d sat til 0 (forsvundet fra listen)"
          % (len(merged), sum(merged.values()), zeroed))

    out_dir = args.out or work
    os.makedirs(out_dir, exist_ok=True)
    files = []
    stamp = NOW.strftime("%Y-%m-%d")      # én arkivkopi pr. dag, overskrives i løbet af dagen
    for o in cfg["outputs"]:
        p = os.path.join(out_dir, o["file"])
        write_xlsx(p, o["column"], merged)
        files.append((p, target["dir"] + "/" + o["file"]))
        files.append((p, target["dir"] + "/arkiv/" + stamp + "_" + o["file"]))
        vprint("  skrevet", p, "->", o["column"])

    state_path = os.path.join(out_dir, cfg["state_file"])
    json.dump(state, open(state_path, "w"), indent=1)

    if local_mode:
        print("\nLokal test - intet uploadet.")
        return
    files.append((state_path, state_remote))
    upload(files, target)
    print("OK - fjernlager publiceret for %d kilde(r)" % len(per_supplier))


if __name__ == "__main__":
    main()
